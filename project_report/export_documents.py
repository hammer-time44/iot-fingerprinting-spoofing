import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

def clean_markdown_inline(text):
    text = text.replace('&', '&amp;')
    text = text.replace('<', '&lt;').replace('>', '&gt;')
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'(?<!\*)\*([^\*]+?)\*(?!\*)', r'<i>\1</i>', text)
    text = re.sub(r'`(.*?)`', r'<font name="Courier">\1</font>', text)
    return text

def export_docx(md_path, docx_path, figures_dir):
    print("Generating DOCX...")
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    
    for line in lines:
        raw_line = line.rstrip('\n')
        
        if raw_line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue
            
        if in_code_block:
            code_lines.append(raw_line)
            continue
            
        if raw_line.strip().startswith('|') and raw_line.strip().endswith('|'):
            cells = [c.strip() for c in raw_line.strip()[1:-1].split('|')]
            if all(set(c) <= set('-: ') for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        else:
            if in_table and table_rows:
                ncols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=ncols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.style = 'Table Grid'
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < ncols:
                            cell = t.cell(r_idx, c_idx)
                            cell.text = cell_text.replace('**', '').replace('*', '')
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.bold = True
                                        r.font.size = Pt(8.5)
                            else:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.size = Pt(8)
                doc.add_paragraph()
                table_rows = []
                in_table = False
                
        if not raw_line.strip():
            continue
            
        if raw_line.strip() == '---':
            p = doc.add_paragraph()
            p.add_run('_______________________________________________________________________________')
            continue
            
        if raw_line.startswith('# '):
            h = doc.add_heading(raw_line[2:], level=1)
            h.runs[0].font.color.rgb = RGBColor(0x1a, 0x36, 0x5d)
            continue
        elif raw_line.startswith('## '):
            h = doc.add_heading(raw_line[3:], level=2)
            h.runs[0].font.color.rgb = RGBColor(0x2b, 0x6c, 0xb0)
            continue
        elif raw_line.startswith('### '):
            h = doc.add_heading(raw_line[4:], level=3)
            h.runs[0].font.color.rgb = RGBColor(0x2c, 0x52, 0x82)
            continue
        elif raw_line.startswith('#### '):
            h = doc.add_heading(raw_line[5:], level=4)
            continue
            
        if raw_line.strip().startswith('* ') or raw_line.strip().startswith('- '):
            text = raw_line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            continue
            
        m = re.match(r'^\s*(\d+)\.\s+(.*)', raw_line)
        if m:
            num, text = m.groups()
            p = doc.add_paragraph(style='List Number')
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            continue
            
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', raw_line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
                
    doc.add_page_break()
    doc.add_heading('Visual Results & Performance Figures', level=1)
    
    figures = [
        ('fig1_confusion_matrix.png', 'Figure 1: Closed-Set Test Confusion Matrix (N=1,067, Accuracy: 99.34%) [results/step7_20261026_151020/cm_test.png]'),
        ('fig2_per_device_f1.png', 'Figure 2: Per-Device F1-Score Comparison: DWT (93.98%) vs Raw Statistical Moments (99.38%) [results/step7_20260917_144006 vs step7_20261026_151020]'),
        ('fig3_feature_importance.png', 'Figure 3: Random Forest Gini Feature Importances for 13 Timing Moments [results/step7_20261026_151020/rf_model.pkl]'),
        ('fig4_validation_confidence_hist.png', 'Figure 4: Validation Set Confidence Distribution & Calibrated Rejection Threshold tau=0.84 [results/step7_20261026_151020/results_summary.json]'),
        ('fig5_unknown_device_detection.png', 'Figure 5: Empirical Open-Set Rejection Rates Across 15 Foreign Hardware Device Types [results/unknown_evaluation_table.csv]'),
        ('fig6_spoof_rejection_by_experiment.png', 'Figure 6: Anti-Spoofing & Impersonation Rejection Success Rates Across Scenarios [results/spoof_20261028_011238/spoof_results.json]'),
        ('fig7_mimicry_jitter_vs_strength.png', 'Figure 7: Detection Robustness vs. Adversarial Mimicry and Jitter Strength [results/spoof_20261028_011238/spoof_results.json]')
    ]
    
    for fname, caption in figures:
        fpath = os.path.join(figures_dir, fname)
        if os.path.exists(fpath):
            doc.add_paragraph()
            p_cap = doc.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.bold = True
            r_cap.font.size = Pt(9.5)
            doc.add_picture(fpath, width=Inches(5.8))
            doc.add_paragraph()
            
    doc.save(docx_path)
    print(f"DOCX saved to: {docx_path}")

def export_pdf(md_path, pdf_path, figures_dir):
    print("Generating PDF via ReportLab...")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=0.5*inch,
        rightMargin=0.5*inch,
        topMargin=0.5*inch,
        bottomMargin=0.5*inch
    )
    
    styles = getSampleStyleSheet()
    
    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#222222')
    )
    
    h1_style = ParagraphStyle(
        'CustomH1',
        parent=styles['Heading1'],
        fontSize=14,
        leading=17,
        textColor=colors.HexColor('#1a365d'),
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'CustomH2',
        parent=styles['Heading2'],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#2b6cb0'),
        spaceBefore=9,
        spaceAfter=4,
        keepWithNext=True
    )
    
    h3_style = ParagraphStyle(
        'CustomH3',
        parent=styles['Heading3'],
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor('#2c5282'),
        spaceBefore=7,
        spaceAfter=3,
        keepWithNext=True
    )
    
    bullet_style = ParagraphStyle(
        'CustomBullet',
        parent=styles['Normal'],
        fontSize=8,
        leading=10.5,
        leftIndent=15,
        firstLineIndent=-10,
        textColor=colors.HexColor('#222222'),
        spaceAfter=2
    )
    
    code_style = ParagraphStyle(
        'CustomCode',
        parent=styles['Code'],
        fontSize=7,
        leading=8.5,
        fontName='Courier',
        textColor=colors.HexColor('#2d3748'),
        backColor=colors.HexColor('#f7fafc'),
        borderColor=colors.HexColor('#e2e8f0'),
        borderWidth=0.5,
        borderPadding=3,
        spaceBefore=3,
        spaceAfter=5
    )
    
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontSize=7,
        leading=8.5,
        textColor=colors.HexColor('#222222')
    )
    
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontSize=7,
        leading=8.5,
        fontName='Helvetica-Bold',
        textColor=colors.white
    )

    story = []
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    
    for line in lines:
        raw_line = line.rstrip('\n')
        
        if raw_line.startswith('```'):
            if in_code_block:
                in_code_block = False
                clean_code = '\n'.join(code_lines).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(f"<pre>{clean_code}</pre>", code_style))
                story.append(Spacer(1, 3))
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue
            
        if in_code_block:
            code_lines.append(raw_line)
            continue
            
        if raw_line.strip().startswith('|') and raw_line.strip().endswith('|'):
            cells = [c.strip() for c in raw_line.strip()[1:-1].split('|')]
            if all(set(c) <= set('-: ') for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
            continue
        else:
            if in_table and table_rows:
                ncols = max(len(r) for r in table_rows)
                formatted_table_data = []
                for r_idx, row in enumerate(table_rows):
                    row_cells = []
                    for c_idx in range(ncols):
                        cell_txt = row[c_idx] if c_idx < len(row) else ''
                        clean_txt = clean_markdown_inline(cell_txt)
                        if r_idx == 0:
                            p = Paragraph(clean_txt, table_header_style)
                        else:
                            p = Paragraph(clean_txt, table_cell_style)
                        row_cells.append(p)
                    formatted_table_data.append(row_cells)
                    
                col_width = (7.5 * inch) / ncols
                t = Table(formatted_table_data, colWidths=[col_width]*ncols)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2b6cb0')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e0')),
                    ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#a0aec0')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7fafc')])
                ]))
                story.append(t)
                story.append(Spacer(1, 5))
                table_rows = []
                in_table = False
                
        if not raw_line.strip():
            continue
            
        if raw_line.strip() == '---':
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0'), spaceAfter=6, spaceBefore=6))
            continue
            
        clean_line = clean_markdown_inline(raw_line)
        
        try:
            if raw_line.startswith('# '):
                story.append(Paragraph(clean_line[2:], h1_style))
            elif raw_line.startswith('## '):
                story.append(Paragraph(clean_line[3:], h2_style))
            elif raw_line.startswith('### '):
                story.append(Paragraph(clean_line[4:], h3_style))
            elif raw_line.startswith('#### '):
                story.append(Paragraph(clean_line[5:], h3_style))
            elif raw_line.strip().startswith('* ') or raw_line.strip().startswith('- '):
                story.append(Paragraph(f"&bull; {clean_line.strip()[2:]}", bullet_style))
            elif re.match(r'^\s*\d+\.\s+', raw_line):
                m = re.match(r'^\s*(\d+)\.\s+(.*)', clean_line)
                story.append(Paragraph(f"<b>{m.group(1)}.</b> {m.group(2)}", bullet_style))
            else:
                story.append(Paragraph(clean_line, normal_style))
                story.append(Spacer(1, 3))
        except Exception:
            # Fallback to plain text if XML parsing trips on unexpected symbol
            plain = raw_line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(plain, normal_style))
            
    story.append(PageBreak())
    story.append(Paragraph("Visual Evidence Figures", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#2b6cb0'), spaceAfter=8))
    
    figures = [
        ('fig1_confusion_matrix.png', 'Figure 1: Closed-Set Test Confusion Matrix (N=1,067, Accuracy: 99.34%) [results/step7_20261026_151020/cm_test.png]'),
        ('fig2_per_device_f1.png', 'Figure 2: Per-Device F1-Score Comparison: DWT (93.98%) vs Raw Statistical Moments (99.38%) [results/step7_20260917_144006 vs step7_20261026_151020]'),
        ('fig3_feature_importance.png', 'Figure 3: Random Forest Gini Feature Importances for 13 Timing Moments [results/step7_20261026_151020/rf_model.pkl]'),
        ('fig4_validation_confidence_hist.png', 'Figure 4: Validation Set Confidence Distribution & Calibrated Rejection Threshold tau=0.84 [results/step7_20261026_151020/results_summary.json]'),
        ('fig5_unknown_device_detection.png', 'Figure 5: Empirical Open-Set Rejection Rates Across 15 Foreign Hardware Device Types [results/unknown_evaluation_table.csv]'),
        ('fig6_spoof_rejection_by_experiment.png', 'Figure 6: Anti-Spoofing & Impersonation Rejection Success Rates Across Scenarios [results/spoof_20261028_011238/spoof_results.json]'),
        ('fig7_mimicry_jitter_vs_strength.png', 'Figure 7: Detection Robustness vs. Adversarial Mimicry and Jitter Strength [results/spoof_20261028_011238/spoof_results.json]')
    ]
    
    for fname, caption in figures:
        fpath = os.path.join(figures_dir, fname)
        if os.path.exists(fpath):
            img = Image(fpath, width=6.2*inch, height=3.9*inch)
            cap = Paragraph(f"<b>{caption}</b>", bullet_style)
            story.append(KeepTogether([img, Spacer(1, 3), cap, Spacer(1, 10)]))
            
    doc.build(story)
    print(f"PDF saved to: {pdf_path}")

if __name__ == '__main__':
    md_file = 'project_report/REPORT.md'
    docx_file = 'project_report/REPORT.docx'
    pdf_file = 'project_report/REPORT.pdf'
    fig_dir = 'project_report/figures'
    
    export_docx(md_file, docx_file, fig_dir)
    export_pdf(md_file, pdf_file, fig_dir)
